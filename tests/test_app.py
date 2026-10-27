"""API and AI-layer tests. OpenAI is replaced by a fake client, so nothing leaves the machine."""
import io
import unittest
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace

import docx
from fastapi.testclient import TestClient

from jobhunt import ai, api, db, profile, refresh
from jobhunt.models import Job

RESUME = ["Priya Sharma", "Software Engineer", "Pune, India | priya@example.com",
          "EDUCATION", "Example Institute of Technology", "B.Tech in Computer Science   Aug 2022 – May 2026",
          "SKILLS", "Languages: Java, Python, SQL", "Frameworks: Spring Boot, Docker, AWS, PyTorch",
          "PROJECTS", "Payments API", "• Built REST APIs in Spring Boot with Docker on AWS"]


def resume_docx():
    d = docx.Document()
    for line in RESUME:
        d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def seed_jobs():
    now = datetime.now(timezone.utc)
    jobs = []
    for i, (title, city) in enumerate([("Backend Engineer, New Grad", "Pune, India"),
                                       ("ML Engineer Intern", "Bengaluru, India")]):
        j = Job("Greenhouse", "Acme", title, city, f"https://example.com/{i}", now - timedelta(days=i),
                "Java, Spring Boot, Docker. 0-2 years of experience.")
        j.kind, j.level, j.key, j.score, j.cities = ("SDE" if i == 0 else "ML"), "New grad", f"acme|{i}", 60.0, [city.split(",")[0]]
        jobs.append(j)
    db.upsert_jobs(jobs, now.isoformat(timespec="seconds"))


class FakeCompletions:
    def __init__(self, answer):
        self.answer, self.calls = answer, []

    def parse(self, **kw):
        self.calls.append(kw)
        parsed = kw["response_format"](**self.answer)
        msg = SimpleNamespace(parsed=parsed, refusal=None)
        return SimpleNamespace(choices=[SimpleNamespace(message=msg)],
                               usage=SimpleNamespace(prompt_tokens=1000, completion_tokens=200))


def fake_ai(answer):
    client = ai.AI.__new__(ai.AI)
    client.model, client.effort, client.price = "gpt-6-luna", "low", [0.10, 0.50]
    client.tokens_in = client.tokens_out = 0
    import threading
    client._lock = threading.Lock()
    client.client = SimpleNamespace(chat=SimpleNamespace(completions=FakeCompletions(answer)))
    return client


ASSESSMENT = {"years_required": 0, "entry_level_ok": True, "in_india": True, "work_mode": "hybrid", "fit_score": 88,
              "verdict": "strong", "reasons": ["Spring Boot and Docker match"], "missing_skills": ["Kafka"],
              "must_haves": ["Java"], "best_resume": "Priya", "role_match": True,
              "contacts": [{"name": "Ravi", "role": "Recruiter", "email": "ravi@acme.in", "phone": ""}]}


class Api(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.c = TestClient(api.app)
        r = cls.c.post("/api/resumes", files={"file": ("Priya.docx", resume_docx())}, data={"label": "Priya"})
        assert r.status_code == 200, r.text
        seed_jobs()

    def test_state_and_jobs(self):
        st = self.c.get("/api/state").json()
        self.assertEqual(st["profile"]["name"], "Priya Sharma")
        self.assertFalse(st["ai"]["configured"])
        self.assertTrue(any(l["site"] == "LinkedIn" for g in st["links"] for l in g["links"]))
        jobs = self.c.get("/api/jobs").json()["jobs"]
        self.assertEqual({j["title"] for j in jobs}, {"Backend Engineer, New Grad", "ML Engineer Intern"})
        self.assertTrue(all(j["is_new"] for j in jobs))

    def test_status_and_notes(self):
        jid = self.c.get("/api/jobs").json()["jobs"][0]["id"]
        r = self.c.patch(f"/api/jobs/{jid}", json={"status": "applied", "notes": "Referred by Ravi"})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()["status"], "applied")
        self.assertEqual(self.c.patch(f"/api/jobs/{jid}", json={"status": "bogus"}).status_code, 400)
        self.assertEqual(self.c.patch("/api/jobs/99999", json={"status": "saved"}).status_code, 404)

    def test_rejects_bad_upload(self):
        r = self.c.post("/api/resumes", files={"file": ("x.exe", b"MZ")})
        self.assertEqual(r.status_code, 400)

    def test_ai_needs_key(self):
        jid = self.c.get("/api/jobs").json()["jobs"][0]["id"]
        self.assertEqual(self.c.post(f"/api/jobs/{jid}/assess").status_code, 400)
        self.assertEqual(self.c.put("/api/ai/providers/openai", json={"api_key": ""}).status_code, 400)

    def test_profile_preferences(self):
        r = self.c.put("/api/profile", json={"current_city": "Pune", "preferred_cities": ["Pune", "Bengaluru"],
                                             "roles": ["SDE"], "about": "Back from the US"})
        prefs = r.json()["profile"]["preferences"]
        self.assertEqual(prefs["current_city"], "Pune")
        self.assertEqual(prefs["roles"], ["SDE"])
        self.assertIn("Back from the US", profile.candidate_block())
        self.c.put("/api/profile", json={"roles": ["SDE", "ML"]})

    def test_host_guard(self):
        api.app.state.allowed_hosts = {"127.0.0.1:8000"}
        try:
            c = TestClient(api.app, base_url="http://evil.example")
            self.assertEqual(c.get("/api/state").status_code, 403)
            ok = TestClient(api.app, base_url="http://127.0.0.1:8000")
            self.assertEqual(ok.get("/api/health").status_code, 200)
            self.assertEqual(ok.put("/api/profile", json={}, headers={"Origin": "https://evil.example"}).status_code, 403)
        finally:
            api.app.state.allowed_hosts = None


class AILayer(unittest.TestCase):
    def test_assess_parses_and_tracks_cost(self):
        client = fake_ai(ASSESSMENT)
        out = client.assess("CANDIDATE\n...", {"title": "SDE", "company": "Acme", "location": "Pune",
                                               "description": "Java"})
        self.assertEqual(out["fit_score"], 88)
        self.assertEqual(client.tokens_in, 1000)
        self.assertAlmostEqual(client.cost(), (1000 * 0.10 + 200 * 0.50) / 1e6)
        call = client.client.chat.completions.calls[0]
        self.assertEqual(call["reasoning_effort"], "low")
        self.assertIn("CANDIDATE", call["messages"][0]["content"])      # candidate first, so it caches

    def test_rejection_rules(self):
        self.assertFalse(ai.rejected(ASSESSMENT))
        self.assertTrue(ai.rejected(dict(ASSESSMENT, years_required=3)))
        self.assertTrue(ai.rejected(dict(ASSESSMENT, in_india=False)))
        self.assertTrue(ai.rejected(dict(ASSESSMENT, years_required=None, entry_level_ok=False)))

    def test_ai_pass_scores_and_caches(self):
        db.resumes() or Api.setUpClass()
        client = fake_ai(ASSESSMENT)
        cfg = {"ai": {"max_jobs_per_refresh": 10, "concurrency": 2}}
        checked, err = refresh._ai_pass(client, cfg, 2, refresh.Progress())
        self.assertIsNone(err)
        self.assertGreaterEqual(checked, 1)
        jobs = [db.public_job(r) for r in db.list_jobs()]
        self.assertTrue(all(j["ai"] and j["ai"]["fit_score"] == 88 for j in jobs))
        self.assertTrue(all(j["score"] == round(0.7 * 88 + 0.3 * j["keyword_score"], 1) for j in jobs))
        again, _ = refresh._ai_pass(client, cfg, 2, refresh.Progress())
        self.assertEqual(again, 0)                       # already checked against this profile


if __name__ == "__main__":
    unittest.main()


class StaleLock(unittest.TestCase):
    def test_lock_from_dead_process_is_cleared(self):
        lock = refresh._FileLock.path
        lock.parent.mkdir(parents=True, exist_ok=True)
        lock.write_text("999999")                      # no such process
        with refresh._FileLock():
            self.assertTrue(lock.exists())
        self.assertFalse(lock.exists())

    def test_live_lock_blocks(self):
        import os
        lock = refresh._FileLock.path
        lock.write_text(str(os.getppid()))             # a process that is certainly alive
        try:
            with self.assertRaises(refresh.Busy):
                with refresh._FileLock():
                    pass
        finally:
            lock.unlink(missing_ok=True)


class Contacts(unittest.TestCase):
    def test_only_hiring_contacts_from_the_text(self):
        d = ("Send your resume to hr@acme.in or careers@acme.com. For accommodations email talent@acme.com. "
             "Privacy: privacy@acme.com. Recruiter: priya.k@acme.co.in. Questions? no-reply@acme.com")
        self.assertEqual([c["email"] for c in db.posting_contacts(d)], ["hr@acme.in", "careers@acme.com", "priya.k@acme.co.in"])

    def test_ai_contacts_come_first_without_duplicates(self):
        r = {"description": "Mail hr@acme.in"}
        ai_part = {"contacts": [{"name": "Priya", "role": "Recruiter", "email": "hr@acme.in", "phone": ""}]}
        out = db._contacts(r, ai_part)
        self.assertEqual(len(out), 1)
        self.assertEqual(out[0]["name"], "Priya")


class TrackerImport(unittest.TestCase):
    def test_csv_with_loose_headings(self):
        from jobhunt import tracker_import as T
        data = ("Company Name,Position,Applied On,City,Status\nRazorpay,SDE 1,05/10/2026,,Interview scheduled\n"
                "Swiggy,ML Engineer,2026-10-07,Bengaluru,\n,,,,\nZepto,,01/10/2026,,\n").encode()
        entries, problems = T.parse("apps.csv", data)
        self.assertEqual([(e["company"], e["role"], e["date"], e["location"], e["status"]) for e in entries],
                         [("Razorpay", "SDE 1", "2026-10-05", "", "interviewing"),
                          ("Swiggy", "ML Engineer", "2026-10-07", "Bengaluru", "applied")])
        self.assertEqual(len(problems), 1)

    def test_xlsx_and_api(self):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Comp", "Role", "Date"])
        ws.append(["Meesho", "Software Engineer", __import__("datetime").datetime(2026, 10, 3)])
        buf = io.BytesIO(); wb.save(buf)
        c = TestClient(api.app)
        r = c.post("/api/tracker/import", files={"file": ("apps.xlsx", buf.getvalue())})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["added"], 1)
        again = c.post("/api/tracker/import", files={"file": ("apps.xlsx", buf.getvalue())}).json()
        self.assertEqual((again["added"], again["possible_duplicates"]), (1, 1))  # repeats are kept, only flagged
        jobs = [j for j in c.get("/api/jobs").json()["jobs"] if j["company"] == "Meesho"]
        self.assertEqual((jobs[0]["status"], jobs[0]["location"]), ("applied", ""))
        self.assertEqual(c.post("/api/tracker", json={"company": "X", "role": ""}).status_code, 400)


class RejectRule(unittest.TestCase):
    def test_stated_years_win_over_the_ai_level_guess(self):
        two_years = dict(ASSESSMENT, years_required=2, entry_level_ok=False)
        self.assertFalse(ai.rejected(two_years))                       # "2+ years" still fits under 3
        self.assertTrue(ai.rejected(dict(ASSESSMENT, years_required=None, entry_level_ok=False)))
        self.assertTrue(ai.rejected(dict(ASSESSMENT, role_match=False)))


class Duplicates(unittest.TestCase):
    def setUp(self):
        self.c = TestClient(api.app)

    def test_duplicates_allowed_with_hint(self):
        r = self.c.post("/api/tracker", json={"company": "Husky Technologies Pvt. Ltd.", "role": "Data Scientist", "date": "2026-10-20"})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertIsNone(r.json()["possible_duplicate"])
        r = self.c.post("/api/tracker", json={"company": "HUSKY", "role": "data scientist"})
        self.assertEqual(r.status_code, 200)
        self.assertIn("Husky Technologies Pvt. Ltd. – Data Scientist", r.json()["possible_duplicate"])

    def test_import_keeps_every_row(self):
        csv_ = ("Company,Role,Date,Status\nZomato,SDE,2026-10-01,Screening call\nZomato Pvt Ltd,sde,2026-10-02,\nNykaa,ML Engineer,2026-10-02,\n").encode()
        r = self.c.post("/api/tracker/import", files={"file": ("a.csv", csv_)}).json()
        self.assertEqual((r["added"], r["possible_duplicates"]), (3, 1))
        self.assertTrue(any(j["company"] == "Zomato" and j["status"] == "screening" for j in self.c.get("/api/jobs").json()["jobs"]))

    def test_edit_delete_undo(self):
        jid = self.c.post("/api/tracker", json={"company": "EditCo", "role": "SDE"}).json()["id"]
        r = self.c.patch(f"/api/jobs/{jid}", json={"company": "EditCo 2", "title": "SDE II", "location": "Pune", "url": "https://x.example/1",
                                                   "applied_on": "2026-10-01", "status": "screening", "contact": "hr@x.example"})
        self.assertEqual(r.status_code, 200, r.text)
        j = r.json()
        self.assertEqual((j["company"], j["title"], j["location"], j["status"], j["applied_on"]), ("EditCo 2", "SDE II", "Pune", "screening", "2026-10-01"))
        self.assertEqual(self.c.patch(f"/api/jobs/{jid}", json={"company": ""}).status_code, 400)
        d = self.c.post("/api/tracker/delete", json={"ids": [jid]}).json()
        self.assertEqual(d["deleted"], 1)
        ids = lambda: {x["id"] for x in self.c.get("/api/jobs").json()["jobs"]}
        self.assertNotIn(jid, ids())
        self.c.post("/api/tracker/undo", json={"rows": d["undo"]})
        self.assertIn(jid, ids())

    def test_export_format_reimports(self):
        from jobhunt import tracker_import as T
        exported = ("﻿Date Applied,Company,Role,Job Type,Location,Source,Job Link,Status,Last Update,Days Since Applied,"
                    "Next Step,Contact,Notes\r\n2026-10-21,Acme,SDE,Full-time,,Referral,,OA / Test,2026-10-22,6,"
                    "Prepare,Ravi,\"'=not a formula, \"\"quoted\"\"\"\r\n").encode("utf-8")
        (e,), problems = T.parse("x.csv", exported)
        self.assertEqual((e["status"], e["extra"]["app_source"], e["extra"]["last_update"], e["notes"]),
                         ("oa", "Referral", "2026-10-22", '=not a formula, "quoted"'))


class Backups(unittest.TestCase):
    def test_backup_and_restore_round_trip(self):
        from jobhunt import backup
        c = TestClient(api.app)
        c.post("/api/tracker", json={"company": "BackupCo", "role": "SDE"})
        b = c.post("/api/backups").json()
        self.assertGreaterEqual(b["counts"]["tracked"], 1)
        jid = next(j["id"] for j in c.get("/api/jobs").json()["jobs"] if j["company"] == "BackupCo")
        c.patch(f"/api/jobs/{jid}", json={"status": ""})                       # "lose" it
        self.assertFalse(any(j["company"] == "BackupCo" and j["status"] for j in c.get("/api/jobs").json()["jobs"]))
        r = c.post(f"/api/backups/{b['name']}/restore")
        self.assertEqual(r.status_code, 200, r.text)
        self.assertTrue(any(j["company"] == "BackupCo" and j["status"] == "applied" for j in c.get("/api/jobs").json()["jobs"]))
        kinds = {x["kind"] for x in c.get("/api/backups").json()["backups"]}
        self.assertIn("pre-restore", kinds)                                   # undo point made first
        self.assertEqual(c.post("/api/backups/../../x/restore").status_code, 404)
        bad = c.post("/api/backups/restore-upload", files={"file": ("x.zip", b"not a zip")})
        self.assertEqual(bad.status_code, 400)
        self.assertIsNotNone(backup.auto_daily())
        self.assertIsNone(backup.auto_daily())                                # once a day


class LLMProviders(unittest.TestCase):
    def setUp(self):
        from unittest import mock
        from jobhunt import llm, settings
        self.llm, self.settings = llm, settings
        self.patches = [mock.patch.object(llm, "ollama_models", return_value=[])]
        for p in self.patches:
            p.start()
        self.saved = settings.read_env()
        settings.write_env({k: None for k in llm.KEY_VARS})
        self.c = TestClient(api.app)

    def tearDown(self):
        self.settings.write_env({k: self.saved.get(k) for k in self.llm.KEY_VARS})
        for p in self.patches:
            p.stop()

    def test_no_key_asks_for_one(self):
        self.assertFalse(self.c.get("/api/state").json()["ai"]["configured"])
        r = self.c.post("/api/ai/chat", json={"message": "hi"})
        self.assertEqual(r.status_code, 400)
        self.assertIn("Please add an LLM key", r.json()["detail"])
        self.assertEqual(len(self.c.get("/api/ai/providers").json()["providers"]), len(self.llm.PROVIDERS))
        self.assertIsNone(ai.client_from_settings())

    def test_fallback_order_and_chat_history(self):
        from unittest import mock
        self.settings.write_env({"GROQ_API_KEY": "gsk_test_1234567890", "GEMINI_API_KEY": "g_test_1234567890"})
        self.c.put("/api/ai/order", json={"order": ["groq", "gemini"]})
        self.assertEqual(self.c.get("/api/state").json()["ai"]["active"][:2], ["groq", "gemini"])
        calls = []

        def fake(pid, messages, max_tokens, env=None, model=None):
            calls.append(pid)
            if pid == "groq":
                raise self.llm.LLMError("quota")
            return "Apply to 10 jobs."
        with mock.patch.object(self.llm, "_complete", side_effect=fake):
            self.c.delete("/api/ai/chat")
            r = self.c.post("/api/ai/chat", json={"message": "tips?"}).json()
        self.assertEqual((calls, r["reply"], r["provider"]), (["groq", "gemini"], "Apply to 10 jobs.", "Google Gemini"))
        self.assertEqual([m["role"] for m in self.c.get("/api/ai/chat").json()["messages"]], ["user", "assistant"])
        client = ai.client_from_settings()
        self.assertEqual((client.label, client.model), ("Groq (Llama, fast)", "llama-3.3-70b-versatile"))
        self.c.delete("/api/ai/providers/groq")
        self.assertEqual(self.c.get("/api/state").json()["ai"]["active"], ["gemini"])

    def test_cover_letter_prefers_local_llama(self):
        from unittest import mock
        jid = self.c.post("/api/tracker", json={"company": "LetterCo", "role": "ML Engineer"}).json()["id"]
        db.add_resume("Main", "r.docx", "r.docx", {"text": "Built a RAG search service in Python."}, None)
        seen = []
        with mock.patch.object(self.llm, "ollama_models", return_value=["llama3.2:latest"]), \
                mock.patch.object(self.llm, "_complete", side_effect=lambda pid, *a, **k: seen.append(pid) or "Dear team"):
            r = self.c.post("/api/ai/cover-letter", json={"job_id": jid, "jd": "We need an ML engineer for retrieval systems."})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual((seen, r.json()["provider"], r.json()["company"]), (["ollama"], "Local model (Ollama)", "LetterCo"))

    def test_coach_attachments_and_documents(self):
        from unittest import mock
        self.settings.write_env({"GEMINI_API_KEY": "g_test_1234567890"})
        seen = []
        with mock.patch.object(self.llm, "_complete", side_effect=lambda pid, m, *a, **k: seen.append(m) or "# Asha Rao\n## Skills\n- Python"):
            self.c.delete("/api/ai/chat")
            r = self.c.post("/api/ai/chat", json={"message": "Draft my CV", "files": [{"name": "jd.txt", "text": "Need Python"}],
                                                  "images": [{"name": "hr.png", "data_url": "data:image/png;base64,AAAA"},
                                                             {"name": "bad", "data_url": "javascript:x"}]})
        self.assertEqual(r.status_code, 200, r.text)
        user = seen[0][-1]["content"]
        self.assertEqual([p["type"] for p in user], ["text", "image_url"])           # bad data URL dropped
        self.assertIn("ATTACHED FILE (jd.txt)", user[0]["text"])
        self.assertEqual(r.json()["messages"][0]["attached"], ["jd.txt", "hr.png"])
        for kind, magic in (("pdf", b"%PDF"), ("docx", b"PK")):
            d = self.c.post("/api/ai/document", json={"text": "# Asha Rao\n## Skills\n- Python, “SQL”", "kind": kind, "title": "Asha Rao"})
            self.assertEqual(d.status_code, 200, d.text)
            self.assertTrue(d.content.startswith(magic))
            self.assertIn(f"asha-rao.{kind}", d.headers["content-disposition"])

    def test_context_file_and_pdf(self):
        from unittest import mock
        notes = ("LinkedIn profile. Built a retrieval augmented chatbot used by two hundred students. " * 6).encode()
        x = self.c.post("/api/ai/context-file", files={"file": ("LinkedIn.txt", notes)}).json()
        self.assertEqual(x["name"], "LinkedIn.txt")
        self.assertIn("retrieval augmented chatbot", x["text"])
        self.assertEqual(self.c.post("/api/ai/context-file", files={"file": ("a.exe", b"MZ")}).status_code, 400)
        db.add_resume("Main", "r.docx", "r.docx", {"text": "Python developer."}, None)
        prompts = []
        with mock.patch.object(self.llm, "ollama_models", return_value=["llama3.2"]), \
                mock.patch.object(self.llm, "_complete", side_effect=lambda pid, m, *a, **k: prompts.append(m[1]["content"]) or "Dear team"):
            self.c.post("/api/ai/cover-letter", json={"company": "PdfCo", "role": "SDE", "extra": [x]})
        self.assertIn("EXTRA MATERIAL (LinkedIn.txt)", prompts[0])
        r = self.c.post("/api/ai/cover-letter.pdf", json={"text": "Dear team,\n\nI'd love to join — “quotes” ₹ ok.\n\nBest,\nA",
                                                          "company": "PdfCo", "role": "SDE"})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertTrue(r.content.startswith(b"%PDF"))
        self.assertIn("cover-letter-pdfco.pdf", r.headers["content-disposition"])


class Dashboard(unittest.TestCase):
    def test_stats_math(self):
        from datetime import date
        from jobhunt import stats
        today = date(2026, 10, 27)
        rows = [dict(status="applied", applied_on="2026-10-27", company="Acme", title="ML Engineer", location="Banglore", source="Added by you"),
                dict(status="screening", applied_on="2026-10-27", company="Acme", title="SDE II", location="Hyderabad, India", source="Added by you"),
                dict(status="rejected", applied_on="2026-10-26", company="Beta", title="Data Scientist", location="", source="x"),
                dict(status="offer", applied_on="2026-10-24", company="Gamma", title="Backend Developer", location="Remote", source="x"),
                dict(status="saved", applied_on="", company="Delta", title="SDE", location="", source="x"),
                dict(status="hidden", applied_on="2026-10-27", company="Zed", title="SDE", location="", source="x")]
        s = stats.compute(today, rows)
        self.assertEqual((s["today"]["count"], s["kpis"]["applications"], s["kpis"]["saved"]), (2, 4, 1))
        self.assertEqual(s["kpis"]["streak"], 2)                                   # 26th and 27th
        self.assertEqual(s["kpis"]["response_rate"], 75.0)                         # screening, rejected, offer
        self.assertEqual([f["n"] for f in s["funnel"]], [4, 2, 1, 1, 1])
        self.assertEqual({x["label"] for x in s["locations"]}, {"Bengaluru", "Hyderabad", "Remote", "Not given"})
        self.assertEqual(stats.role_family("Senior ML Engineer"), "ML / AI")
        self.assertEqual(stats.role_family("Backend Developer"), "Software engineering")
        self.assertNotIn("Acme", json_dump(stats.digest(s)))                     # the AI never sees company names

    def test_endpoints(self):
        c = TestClient(api.app)
        self.assertEqual(c.put("/api/settings/target", json={"target": 0}).status_code, 400)
        self.assertEqual(c.put("/api/settings/target", json={"target": 12}).json()["target"], 12)
        self.assertEqual(c.get("/api/state").json()["target"]["target"], 12)
        st = c.get("/api/stats").json()
        self.assertEqual(len(st["per_day"]), 30)
        self.assertEqual(len(st["per_week"]), 12)
        self.assertEqual(st["today"]["target"], 12)
        from unittest import mock
        from jobhunt import llm
        with mock.patch.object(llm, "chat", return_value=("## Summary\n- Keep going", "Test AI")) as m:
            r1 = c.post("/api/ai/insights").json()
            r2 = c.post("/api/ai/insights").json()                                # cached for 12 h
            c.post("/api/ai/insights?force=true")
        self.assertEqual((r1["provider"], r2["at"], m.call_count), ("Test AI", r1["at"], 2))
        self.assertEqual(c.get("/api/ai/insights").json()["text"], "## Summary\n- Keep going")


def json_dump(x):
    import json
    return json.dumps(x)


class Phase4(unittest.TestCase):
    def test_tiers(self):
        from jobhunt import tiers
        cases = {"Google India": "big_tech", "Amazon (AWS)": "big_tech", "Stripe": "us_unicorn", "Observe.AI": "us_unicorn",
                 "Razorpay Software": "startup", "Capgemini": "mid", "X-Rite": "mid", "Meta-Lab": "mid", "Metamorph": "mid"}
        self.assertEqual({k: tiers.tier_of(k) for k in cases}, cases)
        c = TestClient(api.app)
        c.post("/api/tracker", json={"company": "Stripe", "role": "SDE"})
        self.assertIn("US unicorn / scale-up", {t["label"] for t in c.get("/api/stats").json()["tiers"]})

    def test_cold_email_only_to_posting_contacts(self):
        from unittest import mock
        from jobhunt import llm
        jid = db.add_manual("MailCo", "SDE", "2026-10-27")
        with db.conn() as cn:                                   # a posting that prints a hiring address
            cn.execute("UPDATE jobs SET description=? WHERE id=?", ("Apply or write to hiring@mailco.example with your resume.", jid))
        db.add_resume("Main", "r.docx", "r.docx", {"text": "Python developer."}, None)
        c = TestClient(api.app)
        self.assertEqual(c.post("/api/ai/cold-email", json={"job_id": jid, "email": "guess@mailco.example"}).status_code, 400)
        with mock.patch.object(llm, "chat", return_value=("Subject: SDE role\n\nHi", "Test AI")):
            r = c.post("/api/ai/cold-email", json={"job_id": jid, "email": "hiring@mailco.example"})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual((r.json()["to"], r.json()["company"]), ("hiring@mailco.example", "MailCo"))


class Mail(unittest.TestCase):
    def test_classify_and_match(self):
        from jobhunt import mail
        self.assertEqual(mail.classify("Your application", "Unfortunately we will not be moving forward after the interview"), "rejected")
        self.assertEqual(mail.classify("Next steps", "Please complete the HackerRank assessment"), "oa")
        self.assertEqual(mail.classify("Interview invitation", "Please share your availability for"), "interviewing")
        self.assertEqual(mail.classify("Hello", "Monthly newsletter"), None)
        companies = {"stripe": "Stripe", "zomato": "Zomato Pvt Ltd"}
        self.assertEqual(mail.match_company("Stripe Recruiting", "no-reply@greenhouse.io", "Thanks", "", companies), "Stripe")
        self.assertEqual(mail.match_company("Talent", "jobs@zomato.com", "Update", "", companies), "Zomato Pvt Ltd")
        self.assertIsNone(mail.match_company("Talent", "x@acme.com", "Update", "", companies))

    def test_proposals_review_and_send(self):
        from unittest import mock
        from jobhunt import mail
        jid = db.add_manual("Rejecto", "SDE", "2026-10-20")
        msgs = [{"id": "<m1>", "from_name": "Rejecto Careers", "from": "careers@rejecto.example", "subject": "Your application",
                 "date": "2026-10-26T10:00:00+00:00", "text": "Unfortunately we decided not to proceed."},
                {"id": "<m2>", "from_name": "Rejecto Careers", "from": "careers@rejecto.example", "subject": "Thanks for applying",
                 "date": "2026-10-20T10:00:00+00:00", "text": "Thank you for applying"}]
        with mock.patch.object(mail, "fetch", return_value=msgs):
            props = mail.sync()
        mine = [p for p in props if p["job_id"] == jid]
        self.assertEqual([(p["msg_id"], p["proposed"]) for p in mine], [("<m1>", "rejected")])   # confirmation skipped
        self.assertEqual(db.get_job(jid)["status"], "applied")                                    # nothing changed yet
        self.assertEqual(mail.review({"<m1>"}, True), 1)
        j = db.get_job(jid)
        self.assertEqual(j["status"], "rejected")
        self.assertIn("Email: Your application", j["notes"])
        with mock.patch.object(mail, "fetch", return_value=msgs):
            self.assertFalse([p for p in mail.sync() if p["job_id"] == jid])                     # never re-proposed
        c = TestClient(api.app)
        self.assertEqual(c.post("/api/mail/send", json={"to": "a@b.co", "subject": "s", "body": "b"}).status_code, 400)  # not connected
        self.settings_env = __import__("jobhunt.settings", fromlist=["x"])
        self.settings_env.write_env({mail.ADDR_VAR: "me@gmail.com", mail.PASS_VAR: "abcdabcdabcdabcd"})
        try:
            with mock.patch("smtplib.SMTP_SSL") as smtp:
                r = c.post("/api/mail/send", json={"to": "hr@rejecto.example", "subject": "Hello", "body": "Hi", "job_id": jid})
                self.assertEqual(r.status_code, 200, r.text)
                sent = smtp.return_value.__enter__.return_value.send_message.call_args[0][0]
            self.assertEqual((sent["To"], sent["From"], sent["Subject"]), ("hr@rejecto.example", "me@gmail.com", "Hello"))
            self.assertIn("Sent email to hr@rejecto.example", db.get_job(jid)["notes"])
            self.assertEqual(c.post("/api/mail/send", json={"to": "a@b.co, c@d.co", "subject": "s", "body": "b"}).status_code, 400)
        finally:
            self.settings_env.write_env({mail.ADDR_VAR: None, mail.PASS_VAR: None})

    def test_extract_contacts_never_invents_addresses(self):
        from unittest import mock
        from jobhunt import llm
        fake = ('{"contacts": [{"name": "Priya Shah", "email": "priya@acme.example", "title": "Talent Partner", "company": "Acme", '
                '"job_role": "ML Engineer"}, {"name": "Ravi", "email": "ravi@acme.example", "company": "Acme"}, '
                '{"name": "", "email": "not-an-email", "company": "X"}]}')
        c = TestClient(api.app)
        with mock.patch.object(llm, "chat", return_value=(fake, "Test AI")):
            r = c.post("/api/ai/extract-contacts", json={"text": "Hiring ML Engineers! Send CVs to priya@acme.example - Priya, Talent Partner, Acme"}).json()
        got = {x["name"]: x["email"] for x in r["contacts"][:r["found"]]}            # the ones just extracted
        self.assertEqual(got, {"Priya Shah": "priya@acme.example", "Ravi": ""})     # ravi@ wasn't in the text
        db.add_resume("Main", "r.docx", "r.docx", {"text": "Python developer."}, None)
        with mock.patch.object(llm, "chat", return_value=("Subject: ML Engineer - Asha\n\nHi Priya", "Test AI")):
            d = c.post("/api/ai/outreach-email", json={"contact": r["contacts"][0]}).json()
        self.assertTrue(d["text"].startswith("Subject:"))

    def test_contacts_from_postings(self):
        jid = db.add_manual("PostCo", "ML Engineer", "2026-10-27", status="saved")
        with db.conn() as cn:
            cn.execute("UPDATE jobs SET description=? WHERE id=?",
                       ("Send your CV to talent@postco.example. For accommodations email privacy@postco.example.", jid))
        c = TestClient(api.app)
        r = c.post("/api/outreach/from-postings").json()
        mine = [x for x in r["contacts"] if x.get("company") == "PostCo"]
        self.assertEqual([(x["email"], x["job_role"], x["job_id"]) for x in mine], [("talent@postco.example", "ML Engineer", str(jid))])
        self.assertEqual(c.post("/api/outreach/from-postings").json()["added"], 0)          # no duplicates on re-run

    def test_cold_email_counts_as_application(self):
        from jobhunt import stats
        c = TestClient(api.app)
        db.set_setting("outreach_contacts", [{"id": "c1", "name": "Mira", "email": "mira@countco.example", "company": "CountCo",
                                              "job_role": "ML Engineer"}, {"id": "c2", "name": "No Company", "email": "x@y.example"}])
        before = stats.today_count()
        r = c.post("/api/outreach/mark-sent", json={"contact_id": "c1"})
        self.assertEqual(r.status_code, 200, r.text)
        j = db.get_job(r.json()["tracked_job_id"])
        self.assertEqual((j["company"], j["title"], j["status"], j["app_source"]), ("CountCo", "ML Engineer", "applied", "Cold email"))
        self.assertEqual(stats.today_count(), before + 1)
        self.assertTrue(next(x for x in r.json()["contacts"] if x["id"] == "c1")["sent_at"])
        again = c.post("/api/outreach/mark-sent", json={"contact_id": "c1"}).json()
        self.assertEqual(again["tracked_job_id"], j["id"])                        # no duplicate entry
        self.assertEqual(c.post("/api/outreach/mark-sent", json={"contact_id": "c2"}).status_code, 400)

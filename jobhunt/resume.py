"""Resume -> profile: contact details, education, skills, experience, projects, role fit.

Text comes out of PDFs with pypdf's layout mode, which keeps each visual line
intact (plain mode on Google-Docs PDFs yields one word per line). Two-column
lines ("Mphasis          New York, United States") are split on runs of 3+
spaces into fields. Sections are found by their headings; everything else is
regex and the shared skill taxonomy - no network, nothing leaves the machine.
"""
import re
from datetime import date
from pathlib import Path

from .skills import ROLE_OF, find_skills

HEADINGS = {
    "summary": ["summary", "professional summary", "profile", "objective", "career objective", "about me", "about"],
    "education": ["education", "academic background", "academics", "academic qualifications", "qualifications"],
    "skills": ["skills", "technical skills", "core skills", "key skills", "skills and tools", "technologies",
               "tech stack", "skills and interests", "technical expertise", "core competencies"],
    "experience": ["experience", "work experience", "professional experience", "employment", "employment history",
                   "internships", "internship", "internship experience", "work history", "relevant experience"],
    "projects": ["projects", "technical projects", "academic projects", "personal projects", "key projects",
                 "selected projects", "project experience", "major projects"],
    "certifications": ["certifications", "certificates", "courses", "licenses and certifications",
                       "certifications and courses", "online courses"],
    "other": ["achievements", "awards", "honors", "honours", "accomplishments", "publications", "leadership",
              "positions of responsibility", "extracurricular activities", "extra curricular activities",
              "activities", "volunteering", "languages", "interests", "hobbies", "references"],
}
_HEADING_OF = {h: sec for sec, hs in HEADINGS.items() for h in hs}

MONTHS = {m: i for i, ms in enumerate(
    [("jan", "january"), ("feb", "february"), ("mar", "march"), ("apr", "april"), ("may",), ("jun", "june"),
     ("jul", "july"), ("aug", "august"), ("sep", "sept", "september"), ("oct", "october"), ("nov", "november"),
     ("dec", "december")], start=1) for m in ms}
_MON = r"(jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|june?|july?|aug(?:ust)?|sept?(?:ember)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)\.?"
_POINT = rf"(?:{_MON}\s*,?\s*(\d{{4}})|(\d{{1,2}})/(\d{{4}})|(\d{{4}}))"
DATE_RANGE = re.compile(rf"{_POINT}\s*(?:-|–|—|to)\s*(?:{_POINT}|(present|current|now|till date|ongoing))", re.I)

DEGREES = [
    (4, "PhD", re.compile(r"\bph\.?\s?d\b|doctor of philosophy", re.I)),
    (3, "Master's", re.compile(r"\bm\.?\s?s\.?(?=\s|,|\(|$)|\bmasters?\b|master['’]s|\bm\.?\s?tech\b|\bm\.?\s?e\.(?=\s)|"
                              r"\bmca\b|\bmba\b|\bm\.?\s?sc\b", re.I)),
    (2, "Bachelor's", re.compile(r"\bb\.?\s?tech\b|\bb\.?\s?e\.?(?=\s|,|\(|$)|bachelor|\bb\.?\s?sc\b|\bbca\b|"
                                r"\bb\.?\s?s\.?(?=\s|,|\(|$)", re.I)),
    (1, "Diploma", re.compile(r"\bdiploma\b", re.I)),
]

EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
PHONE = re.compile(r"\+?\(?\d[\d\s().-]{8,}\d")
URL = re.compile(r"(?:https?://)?(?:www\.)?(?:[a-z0-9-]+\.)+[a-z]{2,}(?:/[^\s|,)]*)?", re.I)
BULLET = re.compile(r"^\s*[●•▪◦∙·\-–*■►✓➢]\s*")

ML_TITLES = re.compile(r"machine learning|\bml\b|data scien|\bai\b|artificial intelligence|deep learning|\bnlp\b|"
                       r"computer vision|\bllm|generative|applied scien|research scien|data engineer", re.I)
SDE_TITLES = re.compile(r"software|developer|\bsde\b|\bswe\b|full[- ]?stack|back[- ]?end|front[- ]?end|"
                        r"web dev|mobile dev|android|ios dev|devops|programmer", re.I)


class ResumeError(Exception):
    pass


# ---------------------------------------------------------------- text extraction

def extract(path):
    """Return (text, links) for a PDF, DOCX or TXT resume."""
    p = Path(path)
    if not p.exists():
        raise ResumeError(f"Resume not found: {p}")
    try:
        with open(p, "rb") as fh:
            fh.read(1)
    except OSError as e:
        raise ResumeError(f"Can't open {p.name} ({e}). If it's in OneDrive, right-click it and choose "
                          f"'Always keep on this device', or start OneDrive, then run again.") from e
    ext = p.suffix.lower()
    if ext == ".pdf":
        text, links = _pdf(p)
    elif ext == ".docx":
        text, links = _docx(p)
    elif ext in (".txt", ".md"):
        text, links = p.read_text(encoding="utf-8", errors="replace"), []
    else:
        raise ResumeError(f"Unsupported resume type {ext}: use PDF, DOCX or TXT")
    if len(text.split()) < 40:
        raise ResumeError(f"Found almost no text in {p.name}. If it's a scanned image, export it from your "
                          f"editor as a text PDF or DOCX instead.")
    return text, links


def _pdf(p):
    from pypdf import PdfReader
    reader = PdfReader(str(p))
    pages, links = [], []
    for page in reader.pages:
        try:
            t = page.extract_text(extraction_mode="layout")
        except Exception:
            t = page.extract_text() or ""
            # Plain mode on some generators gives one word per line; glue those back together.
            lines = [l for l in t.splitlines() if l.strip()]
            if lines and sum(len(l.split()) for l in lines) / len(lines) < 1.6:
                t = re.sub(r"\s*\n\s*", " ", t)
        pages.append(t)
        for annot in page.get("/Annots") or []:
            try:
                uri = annot.get_object().get("/A", {}).get("/URI")
                if uri:
                    links.append(str(uri))
            except Exception:
                pass
    return "\n".join(pages), links


def _docx(p):
    import docx
    d = docx.Document(str(p))
    lines = [para.text for para in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            lines.append("   ".join(c.text for c in row.cells))
    links = [r.target_ref for r in d.part.rels.values() if "hyperlink" in r.reltype]
    return "\n".join(lines), links


# ---------------------------------------------------------------- parsing helpers

def _fields(line):
    return [f.strip() for f in re.split(r"\s{3,}|\t", line.strip()) if f.strip()]


def _heading(line):
    s = re.sub(r"[^a-z& ]", "", line.lower()).replace("&", "and")
    s = re.sub(r"\s+", " ", s).strip()
    return _HEADING_OF.get(s)


def _month_index(m, y):
    return int(y) * 12 + MONTHS[m.lower()[:3]] - 1


def _point(groups):
    mon, y1, mm, y2, y3 = groups
    if y1:
        return _month_index(mon, y1)
    if y2:
        return int(y2) * 12 + int(mm) - 1
    if y3:
        return int(y3) * 12
    return None


def date_ranges(text, today=None):
    """[(start_index, end_index)] month indexes (year*12 + month-1) for every range in text."""
    today = today or date.today()
    now = today.year * 12 + today.month - 1
    out = []
    for m in DATE_RANGE.finditer(text):
        g = m.groups()
        start = _point(g[0:5])
        end = now if g[10] else _point(g[5:10])
        if start is not None and end is not None and end >= start:
            out.append((start, end))
    return out


def _merged_months(ranges):
    total, last_end = 0, -1
    for s, e in sorted(ranges):
        s = max(s, last_end + 1)
        if e >= s:
            total += e - s + 1
            last_end = e
    return total


def _fmt_month(idx):
    return f"{date(idx // 12, idx % 12 + 1, 1):%b %Y}"


# ---------------------------------------------------------------- profile

def parse(path, today=None):
    text, links = extract(path)
    return build_profile(text, links, label=Path(path).stem, path=str(Path(path).resolve()), today=today)


def build_profile(text, links=(), label="resume", path="", today=None):
    today = today or date.today()
    lines = [l.rstrip() for l in text.splitlines()]
    sections, current = {"header": []}, "header"
    for line in lines:
        if not line.strip():
            continue
        sec = _heading(line) if len(line.split()) <= 6 else None
        if sec:
            current = sec
            sections.setdefault(sec, [])
            continue
        sections.setdefault(current, []).append(line)

    prof = {"label": label, "path": path}
    prof.update(_contact(sections.get("header", []), text, links))
    prof["education"] = _education(sections.get("education", []), today)
    edu_levels = [e["level"] for e in prof["education"]]
    prof["highest_degree"] = max(prof["education"], key=lambda e: e["level"])["degree"] if edu_levels else ""
    prof["degree_level"] = max(edu_levels) if edu_levels else 0
    ends = [e["end"] for e in prof["education"] if e.get("end") is not None]
    grad = max(ends) if ends else None
    prof["graduation"] = _fmt_month(grad) if grad is not None else ""
    prof["graduated"] = grad is None or grad <= today.year * 12 + today.month - 1

    exp_lines = sections.get("experience", [])
    ranges, intern_ranges = [], []
    for i, line in enumerate(exp_lines):
        for r in date_ranges(line, today):
            ranges.append(r)
            ctx = " ".join(exp_lines[max(0, i - 2): i + 1]).lower()
            if "intern" in ctx or "trainee" in ctx:
                intern_ranges.append(r)
    prof["experience_months"] = _merged_months(ranges)
    prof["internship_months"] = _merged_months(intern_ranges)
    prof["fulltime_months"] = _merged_months([r for r in ranges if r not in intern_ranges])

    prof["entries"] = _entries(exp_lines, "experience") + _entries(sections.get("projects", []), "project")
    prof["projects"] = [e["name"] for e in prof["entries"] if e["kind"] == "project"]

    body = "\n".join(l for sec, ls in sections.items() if sec != "header" for l in ls)
    skills = find_skills(body)
    listed = find_skills("\n".join(sections.get("skills", [])))
    used = find_skills("\n".join(exp_lines + sections.get("projects", [])))
    prof["skills"] = sorted(skills, key=lambda s: (-(s in used) - (s in listed), -skills[s], s))
    prof["skills_used"] = sorted(used)
    prof["role_scores"], prof["roles"] = _roles(skills, prof.get("headline", ""), exp_lines, used)
    prof["text"] = re.sub(r"\s+", " ", body)
    return prof


def _contact(header, text, links):
    top = header[:8] or text.splitlines()[:8]
    out = {"name": "", "headline": "", "email": "", "phone": "", "location": "", "linkedin": "", "github": "",
           "portfolio": ""}
    for line in top:
        s = line.strip()
        words = s.split()
        if not out["name"] and 2 <= len(words) <= 5 and re.fullmatch(r"[A-Za-z.' -]+", s) and "resume" not in s.lower():
            out["name"] = s
            continue
        if out["name"] and not out["headline"] and "@" not in s and not re.search(r"\d{3}", s) and len(s) < 110:
            out["headline"] = s
    m = EMAIL.search(" ".join(top)) or EMAIL.search(text)
    out["email"] = m.group(0) if m else ""
    for line in top:
        for seg in re.split(r"\s*[|•·]\s*|\s{3,}", line.strip()):
            digits = re.sub(r"\D", "", seg)
            if not out["phone"] and PHONE.fullmatch(seg.strip()) and 10 <= len(digits) <= 13:
                out["phone"] = re.sub(r"\s*-\s*", "-", seg.strip())
            elif not out["location"] and re.fullmatch(r"[A-Z][A-Za-z .]+,\s*[A-Z][A-Za-z .]+", seg.strip()) \
                    and "@" not in seg and len(seg) < 40:
                out["location"] = seg.strip()
    urls = list(links) + [m.group(0) for m in URL.finditer(" ".join(top)) if "@" not in m.group(0)]
    for u in urls:
        low = u.lower()
        if low.startswith("mailto:") or low.startswith("tel:"):
            continue
        full = u if low.startswith("http") else "https://" + u
        if "linkedin.com/" in low and not out["linkedin"]:
            out["linkedin"] = full
        elif "github.com/" in low and not out["github"]:
            out["github"] = full
        elif not out["portfolio"] and re.search(r"\.(?:dev|io|me|app|com|in|site|page|xyz)(?:/|$)", low) \
                and not any(x in low for x in ("linkedin", "github", "gmail", "leetcode", "google.com")):
            out["portfolio"] = full
    return out


def _education(lines, today):
    out = []
    for i, line in enumerate(lines):
        for level, name, rx in DEGREES:
            if rx.search(line):
                fields = _fields(line)
                field = ""
                m = re.search(r"\b(?:in|of)\s+([A-Za-z&,. ]+?)(?:\s*[|(,]|\s{2,}|$)", fields[0])
                if m:
                    field = m.group(1).strip()
                window = " ".join(lines[i:i + 2])
                rng = date_ranges(window, today)
                prev = _fields(lines[i - 1])[0] if i > 0 else ""
                school = prev if re.search(r"universit|college|institute|school|iit|nit|iiit", prev, re.I) else ""
                out.append({"level": level, "degree": name, "field": field, "school": school,
                            "end": rng[0][1] if rng else None})
                break
    return out


def _entries(lines, kind):
    blocks, cur = [], None
    for line in lines:
        if BULLET.match(line):
            if cur is None:
                cur = {"head": [], "body": []}
                blocks.append(cur)
            cur["body"].append(BULLET.sub("", line).strip())
        elif cur is not None and not cur["body"] and len(cur["head"]) < 3:
            cur["head"].append(line.strip())
        elif cur is not None and cur["body"] and not _fields(line)[0][:1].isupper():
            cur["body"][-1] += " " + line.strip()        # wrapped bullet line
        else:
            cur = {"head": [line.strip()], "body": []}
            blocks.append(cur)
    out = []
    for b in blocks:
        if not b["head"]:
            continue
        f1 = _fields(b["head"][0])
        name = f1[0]
        if kind == "experience" and len(b["head"]) > 1:
            role = _fields(b["head"][1])[0]
            name = f"{role} at {f1[0]}"
        name = DATE_RANGE.sub("", name).strip(" –-|")
        text = " ".join(b["head"] + b["body"])
        out.append({"kind": kind, "name": name, "text": text, "points": b["body"], "skills": sorted(find_skills(text))})
    return out


def _roles(skills, headline, exp_lines, used):
    score = {"SDE": 0.0, "ML": 0.0}
    for s, n in skills.items():
        w = min(n, 3) * (1.5 if s in used else 1.0)
        role = ROLE_OF[s]
        if role == "ml":
            score["ML"] += w
        elif role == "sde":
            score["SDE"] += w
        else:
            score["ML"] += w / 2
            score["SDE"] += w / 2
    titles = [headline] + [_fields(l)[0] for l in exp_lines if not BULLET.match(l)]
    for t in titles:
        if ML_TITLES.search(t):
            score["ML"] += 6
        if SDE_TITLES.search(t):
            score["SDE"] += 6
    top = max(score.values()) or 1
    norm = {k: round(v / top, 2) for k, v in score.items()}
    roles = [k for k, v in sorted(norm.items(), key=lambda kv: -kv[1]) if v >= 0.6]
    return norm, roles


def summary(prof):
    edu = ", ".join(f"{e['degree']}{' in ' + e['field'] if e['field'] else ''}" for e in prof["education"])
    exp = prof["experience_months"]
    return (f"{prof['name'] or prof['label']}  |  {' + '.join(prof['roles'])} profile  |  {edu or 'education not found'}"
            f"{'  |  graduating ' + prof['graduation'] if not prof['graduated'] else ('  |  graduated ' + prof['graduation'] if prof['graduation'] else '')}"
            f"\n  experience: {exp} months total ({prof['internship_months']} internship)"
            f"\n  {len(prof['skills'])} skills: {', '.join(prof['skills'][:18])}{' ...' if len(prof['skills']) > 18 else ''}")
